import openpyxl
import requests
import time

EXCEL_FILE = "input.xlsx"
API_URL = "http://127.0.0.1:8000/api/excel/update"

last_values = {}

while True:

    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    ws = wb["Sheet1"]

    for row in range(2, ws.max_row + 1):

        cid = ws.cell(row, 1).value
        budget = ws.cell(row, 3).value

        if cid is None or budget is None:
            continue

        current = (cid, budget)

        # Only send if the value changed
        if last_values.get(row) != current:

            response = requests.post(
                API_URL,
                json={
                    "cid": str(cid),
                    "budget": float(budget)
                }
            )

            print(response.json())

            last_values[row] = current

    time.sleep(2)